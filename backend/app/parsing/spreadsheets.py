"""XLSX / XLS / CSV ledgers -> LedgerRow list. No LLM reads the rows.

Every sheet is read; empty or junk sheets are skipped. The header row is found
by scoring the first rows for known column names (English, French, German,
Spanish), so it does not have to be row 1.
"""
from __future__ import annotations

import csv
import datetime as dt
import io
import os
import re
import unicodedata
from decimal import Decimal, InvalidOperation
from typing import Any, Iterable

from app.contracts import ErrorCode, LedgerRow, PipelineError
from app.extraction.normalize import detect_currency, parse_date, parse_decimal

_HEADER_SCAN_ROWS = 30
_MIN_HEADER_SCORE = 2

# Synonyms are compared after lowercasing, removing accents and punctuation.
SYNONYMS: dict[str, list[str]] = {
    "invoice_number": ["invoice no", "invoice number", "invoice", "invoice id", "inv no", "inv", "bill no", "bill",
                       "bill number", "document no", "doc no", "reference", "ref", "facture", "no facture",
                       "numero facture", "n facture", "rechnung", "rechnungsnummer", "rechnung nr", "rechnungsnr",
                       "factura", "numero factura", "n factura", "receipt no", "receipt number"],
    "invoice_date": ["invoice date", "date", "bill date", "issue date", "date facture", "datum", "rechnungsdatum",
                     "fecha", "fecha factura", "document date"],
    "supplier_name": ["supplier", "vendor", "supplier name", "vendor name", "seller", "fournisseur", "lieferant",
                      "proveedor", "merchant", "payee", "from"],
    "customer_name": ["customer", "client", "customer name", "buyer", "bill to", "kunde", "cliente", "to"],
    "subtotal": ["subtotal", "sub total", "net", "net amount", "amount excl tax", "montant ht", "total ht", "netto",
                 "nettobetrag", "base imponible", "subtotal amount"],
    "tax": ["tax", "vat", "gst", "sales tax", "tax amount", "vat amount", "tva", "mwst", "ust", "iva", "steuer"],
    "total": ["total", "amount", "total amount", "grand total", "amount due", "invoice total", "gross", "ttc",
              "montant ttc", "total ttc", "brutto", "gesamtbetrag", "betrag", "importe", "importe total", "summe"],
    "currency": ["currency", "ccy", "cur", "devise", "wahrung", "moneda"],
}
_MONEY_FIELDS = ("subtotal", "tax", "total")


def _norm_header(value: Any) -> str:
    text = unicodedata.normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9 ]", " ", text)).strip()


_LOOKUP = {_norm_header(s): field for field, syns in SYNONYMS.items() for s in syns}


def _match_header(cell: Any) -> str | None:
    key = _norm_header(cell)
    if not key:
        return None
    return _LOOKUP.get(key)


def _find_header(rows: list[list[Any]]) -> tuple[int, dict[int, str]] | None:
    best: tuple[int, dict[int, str]] | None = None
    best_score = 0
    for r, row in enumerate(rows[:_HEADER_SCAN_ROWS]):
        mapping: dict[int, str] = {}
        for c, cell in enumerate(row):
            field = _match_header(cell)
            if field and field not in mapping.values():
                mapping[c] = field
        score = len(mapping)
        if score > best_score:
            best, best_score = (r, mapping), score
    if best is None or best_score < _MIN_HEADER_SCORE:
        return None
    if not ({"invoice_number", "total"} & set(best[1].values())):
        return None
    return best


def _money(value: Any) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, Decimal)):
        return f"{Decimal(value):.2f}"
    if isinstance(value, float):
        try:
            return f"{Decimal(repr(value)):.2f}"
        except InvalidOperation:
            return None
    return parse_decimal(str(value))


def _date(value: Any) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, dt.datetime):
        return value.date().isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    return parse_date(str(value)).iso_date


def _text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    return text or None


def _json_safe(value: Any) -> Any:
    if isinstance(value, (dt.date, dt.datetime)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, float):
        return repr(value)
    return value


def _rows_to_ledger(sheet_name: str, rows: list[list[Any]]) -> list[LedgerRow]:
    found = _find_header(rows)
    if found is None:
        return []
    header_row, mapping = found
    headers = [str(c) if c is not None else f"col{i + 1}" for i, c in enumerate(rows[header_row])]
    out: list[LedgerRow] = []
    for r in range(header_row + 1, len(rows)):
        row = rows[r]
        if not any(v not in (None, "") for v in row):
            continue
        values: dict[str, Any] = {}
        for c, field in mapping.items():
            values[field] = row[c] if c < len(row) else None
        rec: dict[str, Any] = {
            "invoice_number": _text(values.get("invoice_number")),
            "invoice_date": _date(values.get("invoice_date")),
            "supplier_name": _text(values.get("supplier_name")),
            "customer_name": _text(values.get("customer_name")),
        }
        for f in _MONEY_FIELDS:
            rec[f] = _money(values.get(f))
        currency = _text(values.get("currency"))
        if currency and len(currency) == 3 and currency.isalpha():
            rec["currency"] = currency.upper()
        else:
            raw_amount = values.get("total")
            detected = detect_currency(str(raw_amount)) if isinstance(raw_amount, str) else None
            rec["currency"] = detected.currency if detected and detected.currency else None
        if not rec["invoice_number"] and not rec["total"]:
            continue  # notes, subtotals, footers
        if (rec["invoice_number"] or "").lower().startswith(("total", "summe", "grand")):
            continue
        raw = {headers[i] if i < len(headers) else f"col{i + 1}": _json_safe(v) for i, v in enumerate(row)}
        out.append(LedgerRow(sheet_name=sheet_name, row_no=r + 1, raw=raw, **rec))
    return out


# ---------------------------------------------------------------------------
# Readers
# ---------------------------------------------------------------------------

def _xlsx_sheets(file_path: str) -> Iterable[tuple[str, list[list[Any]]]]:
    import openpyxl

    # Open via a file handle: openpyxl rejects paths without an .xlsx
    # extension, and stored uploads are saved under random keys.
    with open(file_path, "rb") as fh:
        wb = openpyxl.load_workbook(fh, read_only=True, data_only=True)
        try:
            for ws in wb.worksheets:
                yield ws.title, [list(r) for r in ws.iter_rows(values_only=True)]
        finally:
            wb.close()


def _xls_sheets(file_path: str) -> Iterable[tuple[str, list[list[Any]]]]:
    import xlrd

    book = xlrd.open_workbook(file_path)
    for sheet in book.sheets():
        rows: list[list[Any]] = []
        for r in range(sheet.nrows):
            row: list[Any] = []
            for c in range(sheet.ncols):
                cell = sheet.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    row.append(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode))
                elif cell.ctype == xlrd.XL_CELL_EMPTY:
                    row.append(None)
                else:
                    row.append(cell.value)
            rows.append(row)
        yield sheet.name, rows


def _decode(raw: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1", errors="replace")


def _csv_sheets(file_path: str) -> Iterable[tuple[str, list[list[Any]]]]:
    with open(file_path, "rb") as fh:
        text = _decode(fh.read())
    sample = text[:20_000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        counts = {d: sample.count(d) for d in ",;\t|"}
        delimiter = max(counts, key=counts.get)
    rows = [list(r) for r in csv.reader(io.StringIO(text), delimiter=delimiter)]
    yield os.path.splitext(os.path.basename(file_path))[0] or "csv", rows


def sniff_format(file_path: str) -> str:
    with open(file_path, "rb") as fh:
        head = fh.read(8)
    if head.startswith(b"PK\x03\x04"):
        return "xlsx"
    if head.startswith(b"\xd0\xcf\x11\xe0"):
        return "xls"
    return "csv"


def parse_spreadsheet_rows(file_path: str, warnings: list[str]) -> list[LedgerRow]:
    if os.path.getsize(file_path) == 0:
        raise PipelineError(ErrorCode.FILE_EMPTY, "the file is empty")
    fmt = sniff_format(file_path)
    reader = {"xlsx": _xlsx_sheets, "xls": _xls_sheets, "csv": _csv_sheets}[fmt]
    ledger: list[LedgerRow] = []
    try:
        for name, rows in reader(file_path):
            sheet_rows = _rows_to_ledger(name, rows)
            if not sheet_rows:
                warnings.append(f"sheet '{name}': no invoice table found, skipped")
            ledger.extend(sheet_rows)
    except PipelineError:
        raise
    except Exception as exc:  # noqa: BLE001 - corrupt workbook, bad zip, etc.
        raise PipelineError(ErrorCode.SPREADSHEET_PARSE_FAILED, f"spreadsheet could not be read: {type(exc).__name__}") from exc
    if not ledger:
        raise PipelineError(ErrorCode.SPREADSHEET_PARSE_FAILED, "no invoice rows found in any sheet")
    return ledger
